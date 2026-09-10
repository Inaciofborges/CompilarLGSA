#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script para compilar dados de múltiplos arquivos Excel
Extrai dados de A18:J18 e adiciona colunas de identificação (Well, MD, CPERM_A, CPOR, HgPOR)
"""

import os
import pandas as pd
from pathlib import Path
from openpyxl import load_workbook
import tkinter as tk
from tkinter import filedialog
import sys

def selecionar_pasta():
    """Permite o usuário selecionar a pasta via interface gráfica."""
    root = tk.Tk()
    root.withdraw()
    root.attributes('-topmost', True)
    pasta = filedialog.askdirectory(title="Selecione a pasta com os arquivos Excel")
    root.destroy()
    return pasta

def extrair_dados_arquivo(caminho_arquivo):
    """
    Extrai dados de um arquivo Excel.
    Retorna um dicionário com dados de cada aba (exceto "graph").
    """
    dados_abas = {}

    try:
        wb = load_workbook(caminho_arquivo)
        abas = [aba for aba in wb.sheetnames if aba.lower() != "graph"]

        for aba in abas:
            ws = wb[aba]

            # Extrair valores de identificação
            well = ws['A9'].value if ws['A9'].value else ""
            md = ws['J9'].value if ws['J9'].value else ""
            cperm_a = ws['J10'].value if ws['J10'].value else ""
            cpor = ws['J11'].value if ws['J11'].value else ""
            hgpor = ws['J12'].value if ws['J12'].value else ""

            # Extrair dados de A18 a J18
            dados_linha = []
            for col in range(1, 11):  # A=1 até J=10
                celula = ws.cell(row=18, column=col)
                dados_linha.append(celula.value)

            dados_abas[aba] = {
                'Well': well,
                'MD': md,
                'CPERM_A': cperm_a,
                'CPOR': cpor,
                'HgPOR': hgpor,
                'dados': dados_linha
            }

        wb.close()
        return dados_abas

    except Exception as e:
        print(f"Erro ao processar {caminho_arquivo}: {e}")
        return {}

def compilar_dados(pasta_origem):
    """Compila dados de todos os arquivos Excel na pasta."""
    dados_compilados = []

    # Encontrar todos os arquivos Excel
    arquivos_excel = list(Path(pasta_origem).glob('*.xlsx')) + \
                     list(Path(pasta_origem).glob('*.xls')) + \
                     list(Path(pasta_origem).glob('*.xlsm'))

    if not arquivos_excel:
        print("Nenhum arquivo Excel encontrado na pasta.")
        return None

    print(f"Encontrados {len(arquivos_excel)} arquivo(s) Excel")

    # Processar cada arquivo
    for arquivo in sorted(arquivos_excel):
        print(f"Processando: {arquivo.name}")
        dados_arquivo = extrair_dados_arquivo(str(arquivo))

        for aba, dados in dados_arquivo.items():
            linha_compilada = [
                dados['Well'],
                dados['MD'],
                dados['CPERM_A'],
                dados['CPOR'],
                dados['HgPOR']
            ] + dados['dados']
            dados_compilados.append(linha_compilada)

    return dados_compilados

def criar_arquivo_saida(dados_compilados, pasta_origem):
    """Cria arquivo Excel com dados compilados."""
    if not dados_compilados:
        print("Nenhum dado foi compilado.")
        return

    # Nomes das colunas
    colunas = ['Well', 'MD', 'CPERM_A', 'CPOR', 'HgPOR', 'A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J']

    # Unidades
    unidades = ['', 'm', 'mD', 'v/v', 'v/v', '', '', '', '', '', '', '', '', '', '']

    # Criar DataFrame
    df = pd.DataFrame(dados_compilados, columns=colunas)

    # Inserir linha de unidades no topo
    df_unidades = pd.DataFrame([unidades], columns=colunas)
    df = pd.concat([df_unidades, df], ignore_index=True)

    # Salvar arquivo
    arquivo_saida = os.path.join(pasta_origem, 'dados_compilados.xlsx')

    # Usar openpyxl para melhor controle de formatação
    with pd.ExcelWriter(arquivo_saida, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Compilado', index=False, header=True)

        # Ajustar largura das colunas
        worksheet = writer.sheets['Compilado']
        for column in worksheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            worksheet.column_dimensions[column_letter].width = adjusted_width

    print(f"\n✓ Arquivo criado com sucesso: {arquivo_saida}")
    print(f"  Total de linhas de dados: {len(df) - 1}")

def main():
    print("=" * 60)
    print("Script de Compilação de Dados Excel")
    print("=" * 60)

    # Obter pasta do usuário
    pasta = selecionar_pasta()

    if not pasta:
        print("Nenhuma pasta foi selecionada.")
        return

    print(f"\nPasta selecionada: {pasta}\n")

    # Compilar dados
    dados = compilar_dados(pasta)

    # Criar arquivo de saída
    if dados:
        criar_arquivo_saida(dados, pasta)
    else:
        print("Nenhum dado foi compilado.")

if __name__ == "__main__":
    main()
